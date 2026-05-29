import io
import os
import json
from docx import Document

TEMPLATE = os.path.abspath(os.path.join(
    os.path.dirname(__file__), "..", "..",
    "医院模板", "文件汇总", "2.2院内竞选需求表.docx"
))


def _set_cell_text(table, row_idx, col_idx, text):
    """完全替换单元格文本（使用第一个唯一单元格）"""
    try:
        cell = table.cell(row_idx, col_idx)
        cell.text = str(text) if text is not None else ""
    except Exception:
        pass


def _chk(val, true_val):
    """勾选框：val 匹配 true_val 时返回 ☑，否则 □"""
    return "☑" if str(val) == str(true_val) else "□"


def generate(demand, project):
    doc = Document(TEMPLATE)
    t = doc.tables[0]

    pname = ""
    year = demand.year or ""
    budget = ""
    if project:
        pname = project.name or ""
        if not year:
            year = str(project.year) if project.year else ""
        budget = f"{project.amount:.2f}" if project.amount else ""

    # ── 表头 ─────────────────────────────────────────────────────
    # Row 0: 需求科室 (col 0) | 归口管理科室 (col 10)
    _set_cell_text(t, 0, 0, f"需求科室\n{demand.demand_dept or ''}")
    _set_cell_text(t, 0, 10, f"归口管理科室\n{demand.manage_dept or ''}")

    # Row 1: 项目名称
    _set_cell_text(t, 1, 0, f"项目名称\n{pname}")

    # ── 第一部分 ──────────────────────────────────────────────────
    # Row 3: 第一部分内容
    cat = demand.category or "货物"
    category_str = f"{_chk(cat,'货物')}货物{_chk(cat,'服务')}服务{_chk(cat,'工程')}工程"
    has_c = "是" if demand.has_consultant == 1 else "否"
    part1 = (
        f"1.1采购单位：内江市第一人民医院\n"
        f"1.2所属年度：{year}年\n"
        f"1.3编制单位：内江市第一人民医院\n"
        f"1.4编制时间：{demand.compile_date or ''}\n"
        f"1.5项目所属分类：{category_str}\n"
        f"1.6预算金额（元）：{budget or 'XXXXXXXXX'}\n"
        f"1.7项目概况：{demand.overview or 'XXXXXXXXX'}\n"
        f"1.8本项目是否有为采购项目提供整体设计、规范编制或者项目管理、监理、检测等服务的供应商：{_chk(has_c,'是')}是{_chk(has_c,'否')}否"
        + (f"\n若是，填写：{demand.consultant_note}" if demand.has_consultant == 1 and demand.consultant_note else "")
    )
    _set_cell_text(t, 3, 0, part1)

    # ── 第三部分 ──────────────────────────────────────────────────
    # Row 5: 采购实施计划
    pm = demand.proc_method or "院内竞选"
    pkg = demand.pkg_split or "不分包采购"
    my = "是" if demand.multi_year == 1 else "否"
    part3 = (
        f"3.1采购方式：{_chk(pm,'院内竞选')}院内竞选{_chk(pm,'院内单一来源')}院内单一来源\n"
        f"3.2采购包划分：{_chk(pkg,'不分包采购')}不分包采购  {_chk(pkg,'分包采购')}分包采购\n"
        f"3.3是否属于一签多年项目：{_chk(my,'是')}是{_chk(my,'否')}否"
    )
    _set_cell_text(t, 5, 0, part3)

    # ── 第四部分：标的情况 ─────────────────────────────────────────
    # Row 9: 包1第一条  Row 10: 包1第二条  Row 11: 包2第一条  Row 12: 包2第二条 ...
    try:
        items = json.loads(demand.items_json or "[]")
    except Exception:
        items = []

    # Fill item rows 9~15 (rows exist in template for pkg1-3)
    item_rows = [9, 10, 11, 12, 13, 14, 15]
    for idx, item in enumerate(items[:len(item_rows)]):
        r = item_rows[idx]
        _set_cell_text(t, r, 0, item.get("pkg", ""))
        _set_cell_text(t, r, 1, item.get("no", ""))
        _set_cell_text(t, r, 3, item.get("name", ""))
        _set_cell_text(t, r, 6, str(item.get("unit_price", "")))
        _set_cell_text(t, r, 7, str(item.get("qty", "")))
        _set_cell_text(t, r, 9, item.get("unit", ""))
        _set_cell_text(t, r, 10, str(item.get("amount", "")))

    # ── 第四部分：分包情况 Row 17 ──────────────────────────────────
    try:
        packages = json.loads(demand.packages_json or "[]")
    except Exception:
        packages = []

    if packages:
        pkg_lines = []
        for i, pkg_info in enumerate(packages):
            pkg_name = pkg_info.get("pkg_name", f"合同包{i+1}")
            budget_v = pkg_info.get("budget", "")
            limit_v = pkg_info.get("limit_price", "")
            rm = pkg_info.get("review_method", "综合评分法")
            pm2 = pkg_info.get("price_method", "固定单价")
            aj = pkg_info.get("allow_joint", "否")
            asub = pkg_info.get("allow_subcontract", "否")
            pkg_lines.append(
                f"{pkg_name}\n"
                f"4.2.1预算金额（元）：{budget_v}，最高限价（元）：{limit_v}\n"
                f"4.2.2评审方法：{_chk(rm,'综合评分法')}综合评分法{_chk(rm,'最低评标价法')}最低评标价法{_chk(rm,'性价比法')}性价比法\n"
                f"4.2.3定价方式：{_chk(pm2,'固定单价')}固定单价{_chk(pm2,'固定总价')}固定总价\n"
                f"4.2.4是否支持联合体投标：{aj}\n"
                f"4.2.5是否允许合同分包：{asub}"
            )
        _set_cell_text(t, 17, 0, "\n\n".join(pkg_lines))

    # ── 第五部分：技术要求 Row 19（内容行）─────────────────────────
    # Row 19 is the 须知, Row 20 is tech content area (but template has no row 20 - goes to row 18=第五部分 title, 19=须知+content)
    # Actually row 19 is the content row based on template
    if demand.tech_req:
        # Append to the 须知 row as separate text
        existing = t.rows[19].cells[0].text
        _set_cell_text(t, 19, 0, existing + "\n\n" + demand.tech_req)

    # ── 第六部分：商务要求 Row 21 ────────────────────────────────
    if demand.biz_req:
        _set_cell_text(t, 21, 0, demand.biz_req)

    # ── 第七部分：特殊资格要求 Row 26 ────────────────────────────
    if demand.special_qual:
        _set_cell_text(t, 26, 0, demand.special_qual)

    # ── 第八部分：评审因素 Rows 30~33 ────────────────────────────
    try:
        review_factors = json.loads(demand.review_factors_json or "[]")
    except Exception:
        review_factors = []

    factor_rows = [30, 31, 32, 33]
    for idx, factor in enumerate(review_factors[:len(factor_rows)]):
        r = factor_rows[idx]
        _set_cell_text(t, r, 0, factor.get("factor", ""))
        _set_cell_text(t, r, 2, factor.get("score", ""))
        _set_cell_text(t, r, 4, factor.get("is_objective", ""))
        _set_cell_text(t, r, 8, factor.get("standard", ""))

    # ── 第九部分：合同管理 Row 35 ─────────────────────────────────
    ct = demand.contract_type or ""
    ct_types = ["买卖合同", "租赁合同", "建设工程合同", "技术合同", "委托合同", "物业管理合同", "其他合同"]
    ct_str = "".join(f"{_chk(tp, tp) if tp in ct.split(',') else '□'}{tp}" for tp in ct_types)
    actual = "是" if demand.actual_settlement == 1 else "否"
    pb = "是" if demand.perf_bond == 1 else "否"
    qb = "是" if demand.quality_bond == 1 else "否"

    perf_bond_detail = ""
    if demand.perf_bond == 1:
        perf_bond_detail = f"\n  比例：{demand.perf_bond_ratio or ''}，缴纳方式：{demand.perf_bond_method or ''}，说明：{demand.perf_bond_note or ''}"

    part9 = (
        f"9.1合同类型：{ct_str}\n"
        f"9.2是否为据实结算：{_chk(actual,'是')}是{_chk(actual,'否')}否\n"
        f"9.3合同履行期限：{demand.contract_period or ''}\n"
        f"9.4合同履约地点：{demand.contract_location or ''}\n"
        f"9.5合同支付约定：{demand.payment_terms or ''}\n"
        f"9.6验收交付标准和方法：{demand.acceptance_standard or ''}\n"
        f"9.7质量保修范围和保修期：{demand.warranty or ''}\n"
        f"9.8知识产权归属：{demand.ip_ownership or ''}\n"
        f"9.9成本补偿和风险分担约定：{demand.cost_risk or ''}\n"
        f"9.10违约责任与解决争议的方法：{demand.breach_dispute or ''}\n"
        f"9.11合同其他条款：{demand.contract_other or ''}\n"
        f"9.12履约保证金：{_chk(pb,'是')}是{_chk(pb,'否')}否{perf_bond_detail}\n"
        f"9.13质量保证金：{_chk(qb,'是')}是{_chk(qb,'否')}否"
    )
    _set_cell_text(t, 35, 0, part9)

    # ── 第十部分：履约验收 Row 37 ─────────────────────────────────
    ao = demand.accept_org or "自行验收"
    is_ = "是" if demand.invite_supplier == 1 else "否"
    ie_ = "是" if demand.invite_expert == 1 else "否"
    ise_ = "是" if demand.invite_service == 1 else "否"
    it_ = "是" if demand.invite_third == 1 else "否"

    part10 = (
        f"10.1验收组织方式：{_chk(ao,'自行验收')}自行验收{_chk(ao,'委托采购代理机构验收')}委托采购代理机构验收\n"
        f"10.2是否邀请本项目的其他供应商：{_chk(is_,'是')}是{_chk(is_,'否')}否\n"
        f"10.3是否邀请专家：{_chk(ie_,'是')}是{_chk(ie_,'否')}否\n"
        f"10.4是否邀请服务对象：{_chk(ise_,'是')}是{_chk(ise_,'否')}否\n"
        f"10.5是否邀请第三方检测机构：{_chk(it_,'是')}是{_chk(it_,'否')}否\n"
        f"10.6履约验收程序：{demand.accept_procedure or ''}\n"
        f"10.7履约验收时间：{demand.accept_time or ''}\n"
        f"10.8验收组织的其他事项：{demand.accept_other or ''}\n"
        f"10.9技术履约验收内容：{demand.tech_accept or ''}\n"
        f"10.10商务履约验收内容：{demand.biz_accept or ''}\n"
        f"10.11履约验收标准：{demand.accept_std or ''}\n"
        f"10.12履约验收其他事项：{demand.accept_other2 or ''}"
    )
    _set_cell_text(t, 37, 0, part10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    proj_num = project.number if project else "unknown"
    filename = f"院内竞选需求表_{proj_num}_{demand.status}.docx"
    return buf, filename
